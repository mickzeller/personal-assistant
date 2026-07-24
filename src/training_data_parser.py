"""
Training data parser — turns raw writing samples into fine-tuning JSONL.

Sources: Perplexity exports (.md), real messages (.docx), conversation
logs, and plain markdown (reviews, notes). Output is chat-format JSONL:
system (who to write as) + user (the task) + assistant (my actual writing).

Quality matters a ton more than quantity here.
"""

import json
import re
from pathlib import Path
from dataclasses import dataclass, field, asdict
from typing import Optional

from rich.console import Console
from rich.table import Table

console = Console()

# ──────────────────────────────────────────────────────────────
# Data structures
# ──────────────────────────────────────────────────────────────

# The personal system prompt (with actual style specifics) lives in
# configs/system_prompt.local.txt, gitignored so the style fingerprint
# stays out of the public repo. This fallback keeps a fresh clone runnable.
_DEFAULT_SYSTEM_PROMPT = (
    "You are a writing assistant that matches the user's personal communication style. "
    "Write in their natural voice. Keep messages concise — get to the point without "
    "unnecessary buildup. Use light humor when appropriate but never forced. "
    "Avoid corporate speak, excessive formality, or over-explaining. "
    "Sound like a real person texting a friend, not an AI."
)

_LOCAL_PROMPT_PATH = Path(__file__).parent.parent / "configs" / "system_prompt.local.txt"
SYSTEM_PROMPT = (
    _LOCAL_PROMPT_PATH.read_text(encoding="utf-8").strip()
    if _LOCAL_PROMPT_PATH.exists()
    else _DEFAULT_SYSTEM_PROMPT
)

# Audience-specific context keywords (and their label) load from
# configs/contexts.local.json, also gitignored — personal-life vocabulary
# stays local. Defaults cover a generic community group.
_COMMUNITY_DEFAULTS = {
    "context_label": "community or volunteer context",
    "text_message_keywords": ["volunteer", "community", "committee"],
    "raw_message_keywords": ["volunteer", "community", "charity", "service project", "committee"],
    "topic_keywords": ["volunteer", "community event", "service project"],
    "long_prompt": "Write a thoughtful message for a community group",
    "short_prompt": "Write a warm, brief message to a community group",
}

_LOCAL_CONTEXTS_PATH = Path(__file__).parent.parent / "configs" / "contexts.local.json"
_LOCAL_CONTEXTS = (
    json.loads(_LOCAL_CONTEXTS_PATH.read_text(encoding="utf-8"))
    if _LOCAL_CONTEXTS_PATH.exists()
    else {}
)
COMMUNITY = {**_COMMUNITY_DEFAULTS, **_LOCAL_CONTEXTS}

# Speaker names for two-person conversation logs — the real names stay in
# the local file; a fresh clone gets neutral placeholders
SELF_ALIASES = [a.lower() for a in _LOCAL_CONTEXTS.get("speaker_self_aliases", ["me"])]
OTHER_ALIASES = [a.lower() for a in _LOCAL_CONTEXTS.get("speaker_other_aliases", ["them"])]


@dataclass
class TrainingExample:
    """One training sample for fine-tuning."""
    user_prompt: str
    assistant_response: str        # the ideal output, in my voice
    context: Optional[str] = None  # audience, tone, etc.
    source_file: str = ""
    example_type: str = ""         # "perplexity", "raw_writing", "review", etc.

    def to_chat_format(self) -> dict:
        """Convert to MLX/OpenAI chat fine-tuning format."""
        messages = [{"role": "system", "content": SYSTEM_PROMPT}]

        # Bake context into the user prompt if we have it
        prompt = self.user_prompt
        if self.context:
            prompt = f"Context: {self.context}\n\nTask: {prompt}"

        messages.append({"role": "user", "content": prompt})
        messages.append({"role": "assistant", "content": self.assistant_response})

        return {"messages": messages}


# ──────────────────────────────────────────────────────────────
# Perplexity Markdown Parser
# ──────────────────────────────────────────────────────────────

def _strip_perplexity_cruft(text: str) -> str:
    """Remove Perplexity branding, footnotes, and decorative HTML."""
    text = re.sub(r'<img[^>]*pplx[^>]*/?>', '', text)                            # logo
    text = re.sub(r'<div[^>]*>⁂</div>', '', text)                                # dividers
    text = re.sub(r'<span[^>]*display:\s*none[^>]*>.*?</span>', '', text)        # hidden footnote spans
    text = re.sub(r'\[\^[\w_]+\]', '', text)                                     # footnote refs like [^5_1]
    text = re.sub(r'^\[\^[\w_]+\]:.*$', '', text, flags=re.MULTILINE)            # footnote definitions
    text = re.sub(r'^\*{3}\s*$', '', text, flags=re.MULTILINE)                   # *** dividers
    text = re.sub(r'\n{3,}', '\n\n', text)                                       # extra blank lines
    return text.strip()


def _extract_quoted_text(text: str) -> Optional[str]:
    """
    Extract the suggested/polished text from a Perplexity response.

    Perplexity presents suggestions in a bunch of formats (blockquotes,
    code blocks, bullet lists, bold quotes, standalone quotes), so we
    try them in order of reliability.
    """
    # 1. Blockquotes — most reliable
    blockquote_lines = []
    for line in text.split('\n'):
        if line.strip().startswith('>'):
            blockquote_lines.append(line.strip().lstrip('> ').strip())

    if blockquote_lines:
        return ' '.join(blockquote_lines).strip()

    # 2. Code blocks
    code_blocks = re.findall(r'```\s*\n?(.*?)\n?\s*```', text, re.DOTALL)
    if code_blocks:
        return code_blocks[-1].strip()

    # 3. Bold-quoted text: **"some text here"**
    bold_quoted = re.findall(r'\*\*"(.+?)"\*\*', text, re.DOTALL)
    if bold_quoted:
        return bold_quoted[-1].strip()

    # 4. Bullet lists with labeled options — grab the LAST one,
    #    Perplexity tends to put the best/most casual option last
    bullet_quotes = re.findall(
        r'-\s+\*\*[^*]+\*\*[:\s]+["\u201c](.+?)["\u201d]',
        text, re.DOTALL
    )
    if bullet_quotes:
        return bullet_quotes[-1].strip()

    # 4b. Bullet options with bold labels but no quotes
    bullet_labeled = re.findall(
        r'-\s+\*\*[^*]+?\*\*[:\s]+(.+?)(?:\n|$)',
        text
    )
    if bullet_labeled:
        # Very short matches are probably just labels
        good_matches = [m.strip().strip('"').strip('\u201c\u201d') for m in bullet_labeled if len(m.strip()) > 20]
        if good_matches:
            return good_matches[-1]

    # 5. Numbered options: **1.** or **Option 1** followed by text
    numbered = re.findall(
        r'\*\*(?:Option\s+)?\d+[.)]\*\*\s*(?:\([^)]+\)[:\s]*)?\n?\s*[>"\u201c]?\s*(.+?)(?:\n\n|\n\*\*|$)',
        text, re.DOTALL
    )
    if numbered:
        return numbered[-1].strip().strip('"').strip('\u201c\u201d')

    # 6. Standalone quoted lines that look like messages
    standalone_quotes = re.findall(r'^["\u201c](.{20,}?)["\u201d]\s*$', text, re.MULTILINE)
    if standalone_quotes:
        return standalone_quotes[-1].strip()

    # 7. Paragraph after "Here's a..." / "Here are..." (single-turn files)
    heres_match = re.search(
        r'(?:Here\'s|Here are)[^\n]*:\s*\n\n(.+?)(?:\n\n|$)',
        text, re.DOTALL
    )
    if heres_match:
        candidate = heres_match.group(1).strip()
        # Only use if it looks like a message (not AI commentary)
        if len(candidate) > 15 and not candidate.startswith(('Would you', 'Do you', 'This ')):
            return candidate

    return None


def _is_user_confirmation(text: str) -> bool:
    """Check if a user turn is approving the previous suggestion."""
    # First line only — the user might paste their final version after "This!"
    first_line = text.strip().split('\n')[0].strip().lower().rstrip('!').rstrip('.')
    confirmations = {
        "this", "perfect", "that's perfect", "love it",
        "that works", "send it", "done", "yes", "yeah", "yep",
        "awesome", "great", "that's great", "that's good",
        "i like that", "i like this",
    }
    # "this!" is a confirmation, but "this is to a..." is a new topic
    if first_line.startswith("this") and not first_line.startswith(("this is", "this one")):
        return True
    return first_line in confirmations


def _classify_writing_context(user_prompt: str) -> str:
    """Infer the audience/context from the prompt so training data has context clues."""
    prompt_lower = user_prompt.lower()

    if any(w in prompt_lower for w in ["sms", "text", "message to my"]):
        if any(w in prompt_lower for w in ["friend", "buddy", "bro", "cousin", "best friend"]):
            return "casual text to a close friend"
        if any(w in prompt_lower for w in ["coworker", "mentor", "colleague", "team", "dev group"]):
            return "work message to a colleague"
        if any(w in prompt_lower for w in COMMUNITY["text_message_keywords"]):
            return COMMUNITY["context_label"]
        return "personal text message"

    if any(w in prompt_lower for w in ["email", "reply to"]):
        return "email or formal reply"

    if any(w in prompt_lower for w in ["linkedin", "post"]):
        return "LinkedIn or social media post"

    if any(w in prompt_lower for w in ["review", "google"]):
        return "online review"

    if any(w in prompt_lower for w in ["ticket", "comment", "jira", "pr "]):
        return "work ticket or code review comment"

    if any(w in prompt_lower for w in COMMUNITY["topic_keywords"]):
        return COMMUNITY["context_label"]

    if any(w in prompt_lower for w in ["group sms", "group text", "group chat"]):
        return "group text message"

    return ""


def parse_perplexity_md(file_path: Path) -> list[TrainingExample]:
    """
    Parse a Perplexity markdown export into training examples.

    Each convo alternates user headings and AI responses split by ---
    dividers. We capture the user's INITIAL request as the prompt and
    the LAST approved version as the output. Multi-topic convos become
    separate examples.
    """
    raw = file_path.read_text(encoding="utf-8")
    cleaned = _strip_perplexity_cruft(raw)

    # split into turns on --- dividers
    sections = re.split(r'\n---\n', cleaned)

    examples = []
    current_request = None
    current_context = None
    last_good_response = None
    original_text = None

    for section in sections:
        section = section.strip()
        if not section:
            continue

        # User turns start with a heading (# ...)
        heading_match = re.match(r'^#\s+(.+?)(?:\n|$)', section, re.DOTALL)

        if heading_match:
            user_text = heading_match.group(1).strip()
            ai_response_text = section[heading_match.end():].strip()

            # Check if user is confirming the previous suggestion
            if _is_user_confirmation(user_text) and last_good_response:
                # The user approved! Save this example
                if current_request and last_good_response:
                    examples.append(TrainingExample(
                        user_prompt=current_request,
                        assistant_response=last_good_response,
                        context=current_context,
                        source_file=file_path.name,
                        example_type="perplexity_approved",
                    ))
                # Reset for the next topic
                current_request = None
                last_good_response = None
                original_text = None

                # The AI's response to "This!" might kick off a new topic
                if ai_response_text:
                    quoted = _extract_quoted_text(ai_response_text)
                    if quoted:
                        last_good_response = quoted
                continue

            is_new_request = any(phrase in user_text.lower() for phrase in [
                "can you help", "could you help", "can you make",
                "could you make", "could you clean", "can you clean",
                "help me make", "help me write",
                "how does this comment sound", "how does this sound",
                "how does this message sound", "how's this response",
                "is this sms okay", "is this okay",
                "i'm sending", "i'm writing", "i'm teaching",
                "i'm answering", "i'm thinking", "i'd like",
                "i'm trying to reach", "i'm reaching out",
                "this is to a separate", "i'm trying to",
                "my text is", "my sister", "my brother",
            ])

            # Feedback = refining a PREVIOUS suggestion. "Can you make this
            # sms sound more natural" is a new request; "more natural please"
            # alone is feedback.
            is_feedback = (
                not is_new_request
                and any(phrase in user_text.lower() for phrase in [
                    "more casual", "more natural", "more fun",
                    "more humble", "more engaging", "more kind",
                    "less formal", "keep it", "which would",
                    "sure let's", "i do like", "let's see",
                    "same for this", "casual conversation",
                    "casual please",
                ])
            )

            if is_new_request:
                # Save the previous example if we have one
                if current_request and last_good_response:
                    examples.append(TrainingExample(
                        user_prompt=current_request,
                        assistant_response=last_good_response,
                        context=current_context,
                        source_file=file_path.name,
                        example_type="perplexity_last_version",
                    ))

                current_request = user_text
                current_context = _classify_writing_context(user_text)
                original_text = user_text
                last_good_response = None

            # User sharing their OWN draft for feedback?
            is_user_sharing_draft = any(phrase in user_text.lower() for phrase in [
                "how's this", "how does this", "how about this",
                "what about this", "this?", "this!",
            ])

            if is_user_sharing_draft or (user_text and not is_new_request and not is_feedback):
                # If the AI approved their draft, the draft itself is the gold example
                ai_lower = ai_response_text.lower() if ai_response_text else ""
                ai_approves = any(phrase in ai_lower for phrase in [
                    "that's a great", "that's perfect", "that's really",
                    "you can send that", "totally good to send",
                    "reads really warm", "fantastic text",
                    "really good version", "that's a fantastic",
                    "perfect as is", "really warm and natural",
                    "already sounds", "already has great tone",
                    "great reply", "hits all the right",
                    "good to send as is", "your version is totally",
                    "great — fun, warm", "great—warm, personal",
                    "this reads really",
                ])

                if ai_approves and current_request:
                    # Extract just the message part from the user's heading
                    # Sometimes the heading IS the message, sometimes it has
                    # a preamble like "How's this" followed by the message
                    user_message = user_text
                    for preamble in ["how's this", "how does this sound",
                                     "how about this", "this?"]:
                        if user_message.lower().startswith(preamble):
                            user_message = user_message[len(preamble):].strip()
                            user_message = user_message.lstrip(':?!. \n')
                            break

                    if user_message and len(user_message) > 15:
                        examples.append(TrainingExample(
                            user_prompt=current_request,
                            assistant_response=user_message,
                            context=current_context,
                            source_file=file_path.name,
                            example_type="perplexity_user_approved",
                        ))
                        last_good_response = user_message
                        continue

            quoted = _extract_quoted_text(ai_response_text)
            if quoted:
                last_good_response = quoted

            # Special case: "how does this sound?" where the user's own text
            # sits right after the heading and the AI just approves it
            # that text might be the best training example
            if not last_good_response and not is_new_request:
                # Check if user is asking for feedback on their OWN text
                user_lower = user_text.lower()
                is_asking_feedback = any(phrase in user_lower for phrase in [
                    "how does this", "how's this", "is this okay",
                    "is this sms okay", "how does it sound",
                    "does this sound",
                ])
                if is_asking_feedback and ai_response_text:
                    # The user's original text might be in ai_response_text
                    # (it's actually the paragraph right after the heading)
                    # Use the full heading text as context + find the draft
                    lines = ai_response_text.split('\n\n')
                    if lines and len(lines[0].strip()) > 15:
                        candidate = lines[0].strip()
                        # Make sure it's not AI commentary
                        if not candidate.lower().startswith(('here', 'your', 'that', 'it', 'sure')):
                            last_good_response = candidate

        else:
            # No heading. This is a continuation or AI-only response
            quoted = _extract_quoted_text(section)
            if quoted:
                last_good_response = quoted

    # Last conversation might not have gotten an explicit confirmation
    if current_request and last_good_response:
        examples.append(TrainingExample(
            user_prompt=current_request,
            assistant_response=last_good_response,
            context=current_context,
            source_file=file_path.name,
            example_type="perplexity_last_version",
        ))

    # Special case: no examples found, but the heading itself IS the message
    # (user submitted their draft as the heading and the AI approved it)
    if not examples and sections:
        first = sections[0].strip()
        heading_match = re.match(r'^#\s+(.+?)(?:\n|$)', first, re.DOTALL)
        if heading_match:
            heading_text = heading_match.group(1).strip()
            ai_text = first[heading_match.end():].strip().lower()

            # If AI says it's good/perfect/great, the heading IS the message
            if any(phrase in ai_text for phrase in [
                "perfect as is", "that's perfect", "sounds great",
                "great as is", "already sounds", "good to send",
                "really good", "that's a great",
            ]):
                context = _classify_writing_context(heading_text)
                prompt = _generate_prompt_for_raw_text(heading_text, context) if context else "Write a casual message"
                examples.append(TrainingExample(
                    user_prompt=prompt,
                    assistant_response=heading_text,
                    context=context,
                    source_file=file_path.name,
                    example_type="perplexity_heading_as_message",
                ))

    return examples


# ──────────────────────────────────────────────────────────────
# DOCX Parser (raw personal writing)
# ──────────────────────────────────────────────────────────────

def parse_docx_raw_writing(file_path: Path) -> list[TrainingExample]:
    """
    Parse a DOCX of raw writing (Slack messages, etc.) — pure gold, zero
    AI involvement. We don't have the original prompt, so we reverse-engineer
    a plausible one from the content.
    """
    from docx import Document

    doc = Document(str(file_path))
    full_text = '\n'.join(p.text.strip() for p in doc.paragraphs if p.text.strip())

    if not full_text or len(full_text) < 10:
        return []

    # Try to infer what kind of message this is
    context = _classify_raw_message(full_text)
    prompt = _generate_prompt_for_raw_text(full_text, context)

    return [TrainingExample(
        user_prompt=prompt,
        assistant_response=full_text,
        context=context,
        source_file=file_path.name,
        example_type="raw_writing",
    )]


def _classify_raw_message(text: str) -> str:
    """Infer the context of a raw message from its content."""
    text_lower = text.lower()

    if any(w in text_lower for w in COMMUNITY["raw_message_keywords"]):
        return COMMUNITY["context_label"]

    if any(w in text_lower for w in ["javascript", "next.js", "code", "deploy",
                                      "pr ", "ticket", "bug", "api", "framework",
                                      "programming", "react", "server"]):
        return "technical or work discussion"

    if any(w in text_lower for w in ["interview", "team", "hr ", "coworker",
                                      "slack", "meeting", "sync"]):
        return "workplace communication"

    if any(w in text_lower for w in ["hash table", "data structure", "algorithm",
                                      "linked-list", "array", "abstract"]):
        return "technical writing or explanation"

    if any(w in text_lower for w in ["party", "birthday", "rsvp"]):
        return "social event"

    if any(w in text_lower for w in ["song", "music", "video"]):
        return "sharing media with friends"

    return "casual conversation"


def _generate_prompt_for_raw_text(text: str, context: str) -> str:
    """Reverse-engineer a plausible user prompt for a piece of raw writing."""
    text_lower = text.lower()
    length = len(text)

    if context == COMMUNITY["context_label"]:
        if length > 500:
            return COMMUNITY["long_prompt"]
        return COMMUNITY["short_prompt"]

    if "technical writing" in context:
        return "Explain this technical concept in an accessible, engaging way"

    if "technical" in context or "work" in context:
        if any(w in text_lower for w in ["hey team", "hey guys", "yo!"]):
            return "Write a casual message to my dev team"
        return "Write a work message that's professional but still sounds like me"

    if "workplace" in context:
        return "Write a friendly message to coworkers"

    if length < 100:
        return "Write a short, casual reply"
    elif length < 300:
        return "Write a casual message to a friend"
    else:
        return "Write a longer, thoughtful message"


# ──────────────────────────────────────────────────────────────
# Conversation Log Parser (two-person SMS/Slack threads)
# ──────────────────────────────────────────────────────────────

def _is_speaker_line(line: str) -> tuple[str, bool]:
    """
    Check if a line is a speaker header from a Slack/SMS export —
    "<name> 2:38 PM", "You sent...", "<name>:emoji: 2:15 PM", or a bare
    name on its own line. Actual names come from the speaker aliases in
    configs/contexts.local.json. Returns ("self"/"other", True) or ("", False).
    """
    stripped = line.strip()
    if not stripped:
        return "", False

    stripped_lower = stripped.lower()

    # "You sent ..." = the person who exported the log
    if stripped_lower.startswith("you sent"):
        return "self", True

    for speaker, aliases in (("self", SELF_ALIASES), ("other", OTHER_ALIASES)):
        # Bare alias on its own line
        if stripped_lower in aliases:
            return speaker, True
        # "<alias>  sent ...", "<alias>:emoji: 2:15", "<alias> 2:38 PM"
        for alias in aliases:
            if re.match(rf'^{re.escape(alias)}(?::[\w-]+:)?\s+(?:sent\s|\.?\s*\d)', stripped_lower):
                return speaker, True

    return "", False


def _is_continuation_timestamp(line: str) -> bool:
    """Bare timestamp like "2:23" — a follow-up message from the same speaker."""
    return bool(re.match(r'^\d{1,2}:\d{2}(?:\s*(?:AM|PM|am|pm))?\s*$', line.strip()))


def _is_slack_noise(line: str) -> bool:
    """Slack noise to strip: emoji reactions, thread markers, date headers, etc."""
    stripped = line.strip()
    if not stripped:
        return True

    # Pure Slack emoji (":joy:", ":pray::skin-tone-3:")
    if re.match(r'^(?::\w[\w-]*:)+\s*$', stripped):
        return True

    # Reaction counts (bare numbers)
    if re.match(r'^\d{1,2}$', stripped):
        return True

    # Thread markers
    if re.match(r'^\d+ repl', stripped.lower()) or 'view thread' in stripped.lower():
        return True

    # Date-only headers: "Mon 7:06 PM", "Yesterday at 10:01 AM", "June 7, 2021"
    if re.match(r'^(?:Mon|Tue|Wed|Thu|Fri|Sat|Sun)\s+\d', stripped):
        return True
    if re.match(r'^(?:Yesterday|Today)\s+at\s+\d', stripped):
        return True
    if re.match(r'^(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)\w*\s+\d', stripped):
        return True

    if stripped.lower() == "(edited)":
        return True

    return False


def _clean_slack_text(text: str) -> str:
    """Strip Slack artifacts from message text."""
    text = re.sub(r'\s*\(edited\)\s*$', '', text)
    emoji_map = {
        ":joy:": "", ":laughing:": "", ":pray:": "",
        ":1000:": "", ":point_up_2:": "", ":fire:": "",
        ":heart:": "", ":thumbsup:": "", ":tada:": "",
        ":four_leaf_clover:": "", ":tired_face:": "",
    }
    for slack_emoji, replacement in emoji_map.items():
        text = text.replace(slack_emoji, replacement)
    # Catch any remaining Slack emoji (skin tones etc.)
    text = re.sub(r':[\w-]+:', '', text)
    return text.strip()


def parse_conversation_md(file_path: Path) -> list[TrainingExample]:
    """
    Parse a Slack/SMS conversation log, extracting only my messages.
    Back-to-back messages (continuation timestamps) get merged into one
    example.
    """
    raw = file_path.read_text(encoding="utf-8")
    lines = raw.split('\n')

    messages = []
    current_speaker = None
    current_lines = []

    for line in lines:
        speaker, is_header = _is_speaker_line(line)

        if is_header:
            # Save the previous message if it was mine
            if current_speaker == "self" and current_lines:
                msg = _clean_slack_text('\n'.join(current_lines))
                if msg:
                    messages.append(msg)

            current_speaker = speaker
            current_lines = []
            continue

        # Continuation timestamp — same speaker, keep the rapid-fire style together
        if _is_continuation_timestamp(line) and current_speaker:
            continue

        if _is_slack_noise(line):
            continue

        # Line belongs to whoever spoke last
        if current_speaker and line.strip():
            current_lines.append(line.strip())

    # Don't forget the last message
    if current_speaker == "self" and current_lines:
        msg = _clean_slack_text('\n'.join(current_lines))
        if msg:
            messages.append(msg)

    examples = []
    for msg in messages:
        context = _classify_raw_message(msg)
        prompt = _generate_prompt_for_raw_text(msg, context)

        # These are all casual friend convos
        if context == "casual conversation":
            context = "casual text to a close friend"

        examples.append(TrainingExample(
            user_prompt=prompt,
            assistant_response=msg,
            context=context,
            source_file=file_path.name,
            example_type="conversation_log",
        ))

    return examples


def is_conversation_log(raw: str) -> bool:
    """
    Detect a two-person conversation log — needs at least 3 of my lines
    and 2 of theirs in the first 100 lines to be confident.
    """
    self_count = 0
    other_count = 0
    for line in raw.split('\n')[:100]:
        speaker, is_header = _is_speaker_line(line)
        if is_header:
            if speaker == "self":
                self_count += 1
            elif speaker == "other":
                other_count += 1

    return self_count >= 3 and other_count >= 2


# ──────────────────────────────────────────────────────────────
# Plain Markdown Parser (reviews, notes, etc.)
# ──────────────────────────────────────────────────────────────

def parse_plain_md(file_path: Path) -> list[TrainingExample]:
    """Parse plain markdown (Google reviews, notes, other direct writing)."""
    raw = file_path.read_text(encoding="utf-8")

    # Perplexity exports get handled by parse_perplexity_md
    if "pplx-full-logo" in raw or "perplexity.ai" in raw:
        return []

    sections = re.split(r'\n——\n|\n---\n|\n\n---\n\n', raw)
    examples = []

    for section in sections:
        section = section.strip()
        if not section or len(section) < 20:
            continue

        text_lower = section.lower()

        if any(w in text_lower for w in ["review", "star", "recommend", "rating"]):
            context = "online review"
            prompt = "Write an honest, detailed review"
        elif "google reviews" in text_lower:
            continue  # skip headers
        else:
            context = _classify_raw_message(section)
            prompt = _generate_prompt_for_raw_text(section, context)

        examples.append(TrainingExample(
            user_prompt=prompt,
            assistant_response=section,
            context=context,
            source_file=file_path.name,
            example_type="plain_writing",
        ))

    return examples


# ──────────────────────────────────────────────────────────────
# Main Pipeline
# ──────────────────────────────────────────────────────────────

def parse_training_directory(
    directory: Path,
    output_path: Path = None,
) -> list[TrainingExample]:
    """Route every file to the right parser and write the JSONL dataset."""
    directory = Path(directory)
    if output_path is None:
        output_path = directory.parent.parent / "data" / "training" / "writing_style.jsonl"

    output_path.parent.mkdir(parents=True, exist_ok=True)

    all_examples = []
    stat_results = {"perplexity_md": 0, "docx": 0, "plain_md": 0, "conversation_log": 0, "skipped": 0}

    console.print(f"\n[bold blue]Parsing training data from: {directory}[/bold blue]\n")

    for file_path in sorted(directory.iterdir()):
        if file_path.name.startswith('.'):
            continue

        suffix = file_path.suffix.lower()

        try:
            if suffix == '.docx':
                samples = parse_docx_raw_writing(file_path)
                stat_results["docx"] += len(samples)

            elif suffix == '.md':
                raw = file_path.read_text(encoding="utf-8")
                if "pplx-full-logo" in raw or "perplexity.ai" in raw:
                    samples = parse_perplexity_md(file_path)
                    stat_results["perplexity_md"] += len(samples)
                elif is_conversation_log(raw):
                    samples = parse_conversation_md(file_path)
                    stat_results["conversation_log"] += len(samples)
                else:
                    samples = parse_plain_md(file_path)
                    stat_results["plain_md"] += len(samples)
            else:
                console.print(f"  [yellow]⚠ Skipping unsupported: {file_path.name}[/yellow]")
                stat_results["skipped"] += 1
                continue

            if samples:
                all_examples.extend(samples)
                console.print(
                    f"  [green]✓[/green] {file_path.name} → {len(samples)} example(s)"
                )
            else:
                console.print(
                    f"  [yellow]○[/yellow] {file_path.name} → no examples extracted"
                )

        except Exception as e:
            console.print(f"  [red]✗[/red] {file_path.name}: {e}")

    # Quality filter — under 10 chars isn't a real message, over 4000 is
    # an essay, and "how's this..." means we grabbed the wrong side
    before_filter = len(all_examples)
    all_examples = [
        ex for ex in all_examples
        if 10 <= len(ex.assistant_response) <= 4000
           and not ex.assistant_response.lower().startswith(("how's this", "how does this", "what do you think"))
    ]
    filtered = before_filter - len(all_examples)
    if filtered:
        console.print(f"  [dim]Filtered {filtered} low-quality examples[/dim]")

    if all_examples:
        with open(output_path, 'w', encoding='utf-8') as f:
            for ex in all_examples:
                json.dump(ex.to_chat_format(), f, ensure_ascii=False)
                f.write('\n')

        console.print(f"\n[bold green]Done! Wrote {len(all_examples)} examples[/bold green]")
        console.print(f"  Output: {output_path}")
    else:
        console.print("\n[yellow]No training examples extracted.[/yellow]")

    table = Table(title="Extraction Summary")
    table.add_column("Source Type", style="cyan")
    table.add_column("Examples", justify="right", style="green")

    table.add_row("Perplexity conversations", str(stat_results["perplexity_md"]))
    table.add_row("Raw writing (DOCX)", str(stat_results["docx"]))
    table.add_row("Plain markdown", str(stat_results["plain_md"]))
    table.add_row("Conversation logs", str(stat_results["conversation_log"]))
    table.add_row("Skipped files", str(stat_results["skipped"]))
    table.add_row("[bold]Total[/bold]", f"[bold]{len(all_examples)}[/bold]")

    console.print()
    console.print(table)

    return all_examples


def inspect_examples(example: list[TrainingExample], n: int = 5):
    """Preview a few training examples to sanity-check quality."""
    console.print(f"\n[bold blue]Previewing {min(n, len(example))} examples:[/bold blue]\n")

    for i, ex in enumerate(example[:n]):
        console.print(f"[bold cyan]── Example {i+1} ({ex.example_type}) ──[/bold cyan]")
        console.print(f"[dim]Source: {ex.source_file}[/dim]")
        if ex.context:
            console.print(f"[dim]Context: {ex.context}[/dim]")
        console.print(f"[yellow]User:[/yellow] {ex.user_prompt[:200]}")
        console.print(f"[green]Assistant:[/green] {ex.assistant_response[:300]}")
        console.print()


# ──────────────────────────────────────────────────────────────
# CLI entry point
# ──────────────────────────────────────────────────────────────

if __name__ == "__main__":
    import sys

    training_dir = Path(__file__).parent.parent / "knowledge_base" / "training_data"

    if len(sys.argv) > 1:
        training_dir = Path(sys.argv[1])

    examples = parse_training_directory(training_dir)

    if examples:
        inspect_examples(examples, n=5)

        stats_path = training_dir.parent.parent / "data" / "training" / "stats.json"
        stats = {
            "total_examples": len(examples),
            "by_type": {},
            "avg_response_length": sum(len(e.assistant_response) for e in examples) / len(examples),
        }
        for ex in examples:
            stats["by_type"][ex.example_type] = stats["by_type"].get(ex.example_type, 0) + 1

        with open(stats_path, 'w') as f:
            json.dump(stats, f, indent=2)

        console.print(f"\n[dim]Stats written to {stats_path}[/dim]")
